"""결과를 구글 시트용 xlsx(탭 4개)로 묶는다. Drive 업로드 시 구글 시트로 변환된다."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

FILL = {"정상": None, "관찰": "FFF3C4", "점검 권고(단발 튐)": "FFD8C2", "이상(지속)": "F8C9CF", "정비 필요 시점": "7A0019",
        "정비 필요": "7A0019"}
res = pd.read_csv("결과_시점별_판정.csv", encoding="utf-8-sig")
summ = pd.read_csv("결과_기법별_요약.csv", encoding="utf-8-sig")
dec = pd.read_csv("결과_정비판정.csv", encoding="utf-8-sig")
res = res.drop(columns=["시각", "설비명"]).rename(columns={
    "탐지_z": "Z-점수", "탐지_md": "마할라노비스", "탐지_iso": "IsolationForest", "탐지_lof": "LOF",
    "탐지_ocsvm": "OneClassSVM", "탐지_cusum": "CUSUM", "마할라노비스": "마할라노비스 거리"})
rules = pd.DataFrame([
    ["데이터", "Day5_설비센서_50건_가상.csv (실습용 가상 데이터, 설비 2대 × 20분 간격 25건)"],
    ["기준(학습) 구간", "설비별 오전 08:00~11:40, 12건을 정상으로 보고 평균·표준편차·모델을 학습"],
    ["특징", "진동·온도·압력을 기준 구간 평균/표준편차로 표준화(z)"],
    ["득표", "6기법이 각각 이상(1)/정상(0) 표를 던짐 → 득표(6)"],
    ["0~1표", "정상"], ["2~3표", "관찰"],
    ["4표 이상 1회", "점검 권고(단발 튐) — 직후 정상 복귀"],
    ["4표 이상 2회 연속", "두 번째 시점 = 정비 필요 시점 (지속 추세)"],
    ["주의", "탐지는 고장 확정이 아님. 설비 자동 정지 금지 — 담당자 검토·승인 후 비가동 시간에 정비 일정 확정"],
    ["주의", "One-Class SVM은 학습 표본 12건으로 과민(설비A 11건 탐지) → 단독 판정 금지, 다수결로 보정"],
    ["재현 코드", "https://github.com/hwangdoohwa/auto-ax-project-1029/tree/main/작업물/ml_anomaly_day5"],
], columns=["항목", "내용"])

wb = Workbook(); wb.remove(wb.active)
HEAD = PatternFill("solid", fgColor="0E1F3E")
def sheet(name, df, status_col=None, widths=None):
    ws = wb.create_sheet(name)
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in row])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = HEAD; c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    if status_col:
        ci = list(df.columns).index(status_col)
        for r in ws.iter_rows(min_row=2):
            f = FILL.get(r[ci].value)
            if f:
                for c in r: c.fill = PatternFill("solid", fgColor=f)
                if f == "7A0019":
                    for c in r: c.font = Font(bold=True, color="FFFFFF")
    for i, col in enumerate(df.columns, 1):
        w = (widths or {}).get(col) or min(40, max(8, max(len(str(col)), *(len(str(v)) for v in df[col].fillna(""))) * 1.6))
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws
sheet("정비 판정", dec, "판정", {"근거": 60, "다음 조치": 60})
sheet("기법별 요약", summ, widths={"원리": 50})
sheet("시점별 판정(50건)", res, "판정")
sheet("판정 규칙", rules, widths={"항목": 20, "내용": 100})
wb.save("이상치탐지_정비필요시점.xlsx"); print("ok")
